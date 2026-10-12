#!/usr/bin/env python3
"""
Shared design library untuk semua Day-XX presentations.
Theme: "Modern Cobalt" — VS Code Dark+ inspired, syntax-highlighted code.

Cara pakai di generate_presentation.py:
    import sys, os
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from pptx_lib import *
"""

import re
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

# ═══════════════════════════════════════════════════════════════════
# 🎨 Modern Cobalt — Color Palette (VS Code Dark+ inspired)
# ═══════════════════════════════════════════════════════════════════

# Backgrounds
BG_DEEP    = RGBColor(0x0D, 0x11, 0x17)   # github bg
BG_MAIN    = RGBColor(0x11, 0x17, 0x22)
BG_CARD    = RGBColor(0x17, 0x1F, 0x2E)
BG_CARD_2  = RGBColor(0x1F, 0x29, 0x3C)
BG_CODE    = RGBColor(0x05, 0x09, 0x10)

# Primary brand colors
PRIMARY    = RGBColor(0x58, 0xA6, 0xFF)   # bright blue
PRIMARY_LT = RGBColor(0x79, 0xC0, 0xFF)
CYAN       = RGBColor(0x39, 0xD0, 0xD8)

# Secondary
VIOLET     = RGBColor(0xC9, 0xA0, 0xFF)   # soft purple
PURPLE     = RGBColor(0xA3, 0x7A, 0xF7)
PINK       = RGBColor(0xFF, 0x7B, 0xC7)
MAGENTA    = RGBColor(0xFF, 0x6E, 0xC7)

# Status
EMERALD    = RGBColor(0x3F, 0xD9, 0x7F)
LIME       = RGBColor(0x9C, 0xE5, 0x4F)
AMBER      = RGBColor(0xFF, 0xC8, 0x57)
ORANGE     = RGBColor(0xFF, 0xA6, 0x57)
ROSE       = RGBColor(0xFF, 0x6B, 0x8A)
RED        = RGBColor(0xF8, 0x71, 0x71)

# Neutrals
WHITE      = RGBColor(0xF0, 0xF6, 0xFC)
SILVER     = RGBColor(0xC9, 0xD1, 0xD9)   # default code text
SLATE      = RGBColor(0x8B, 0x94, 0x9E)
MUTED      = RGBColor(0x6E, 0x76, 0x81)
DIM        = RGBColor(0x4D, 0x55, 0x5F)
BORDER     = RGBColor(0x30, 0x36, 0x3D)

# Backwards-compat
BG_DARK = BG_MAIN
ACCENT  = PRIMARY
INDIGO  = PURPLE
GRAY    = SLATE
GREEN   = EMERALD
DARK2   = BG_CARD
DARK3   = BG_CODE

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


# ═══════════════════════════════════════════════════════════════════
# 🛠️ Low-level Helpers
# ═══════════════════════════════════════════════════════════════════

def new_slide(prs, bg=None):
    if bg is None:
        bg = BG_MAIN
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    sb = slide.background
    fill = sb.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def rect(slide, x, y, w, h, color, line_color=None, line_w=0):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w if line_w > 0 else 0.75)
    return shape


def rounded_rect(slide, x, y, w, h, color, line_color=None, line_w=0):
    shape = slide.shapes.add_shape(
        5, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w if line_w > 0 else 0.75)
    return shape


def circle(slide, cx, cy, r, color, line_color=None, line_w=0):
    shape = slide.shapes.add_shape(
        9, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w if line_w > 0 else 1.5)
    return shape


def shadow_rect(slide, x, y, w, h, color, shadow_offset=0.05):
    """Rectangle dengan shadow halus di belakang (efek depth)."""
    rounded_rect(slide, x + shadow_offset, y + shadow_offset, w, h, BG_DEEP)
    return rounded_rect(slide, x, y, w, h, color)


def tb(slide, text, x, y, w, h, size=14, bold=False, color=WHITE,
       align=PP_ALIGN.LEFT, italic=False, font_name=None):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    if font_name:
        run.font.name = font_name
    return txBox


def multi_tb(slide, lines, x, y, w, h, size=11, color=WHITE, line_gap=1.2,
             align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if isinstance(line, tuple):
            text = line[0]
            line_color = line[1] if len(line) > 1 else color
            line_bold = line[2] if len(line) > 2 else False
        else:
            text = line
            line_color = color
            line_bold = False
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = line_color
        run.font.bold = line_bold
    return txBox


# ═══════════════════════════════════════════════════════════════════
# 🌈 C++ Syntax Highlighter
# ═══════════════════════════════════════════════════════════════════

CPP_KEYWORDS = {
    'int', 'long', 'short', 'char', 'bool', 'void', 'auto', 'const',
    'static', 'unsigned', 'signed', 'float', 'double', 'struct', 'class',
    'public', 'private', 'protected', 'virtual', 'override', 'final',
    'template', 'typename', 'typedef', 'using', 'namespace', 'enum',
    'union', 'if', 'else', 'for', 'while', 'do', 'switch', 'case',
    'default', 'break', 'continue', 'return', 'goto', 'try', 'catch',
    'throw', 'new', 'delete', 'true', 'false', 'nullptr', 'NULL',
    'this', 'sizeof', 'operator', 'inline', 'extern', 'volatile',
    'mutable', 'explicit', 'friend', 'constexpr', 'static_cast',
    'dynamic_cast', 'reinterpret_cast', 'const_cast', 'and', 'or', 'not',
}

CPP_TYPES = {
    'string', 'vector', 'pair', 'tuple', 'map', 'set', 'multiset',
    'unordered_map', 'unordered_set', 'multimap', 'queue', 'deque',
    'stack', 'priority_queue', 'list', 'array', 'bitset', 'size_t',
    'iterator', 'function', 'optional', 'variant', 'greater', 'less',
    'll', 'll64', 'pii', 'vi', 'vll',
}

CPP_BUILTINS = {
    'sort', 'stable_sort', 'reverse', 'find', 'find_if', 'count',
    'count_if', 'min', 'max', 'min_element', 'max_element',
    'lower_bound', 'upper_bound', 'binary_search', 'unique', 'merge',
    'accumulate', 'reduce', 'transform', 'fill', 'fill_n', 'iota',
    'partition', 'rotate', 'next_permutation', 'prev_permutation',
    'all_of', 'any_of', 'none_of', 'for_each', 'gcd', 'lcm', 'abs',
    'push_back', 'pop_back', 'push_front', 'pop_front', 'emplace_back',
    'emplace', 'insert', 'erase', 'clear', 'size', 'empty', 'front',
    'back', 'begin', 'end', 'rbegin', 'rend', 'top', 'push', 'pop',
    'first', 'second', 'make_pair', 'make_tuple', 'tie', 'get', 'swap',
    'move', 'to_string', 'stoi', 'stoll', 'stod', 'substr', 'length',
    'at', 'data', 'cin', 'cout', 'cerr', 'endl', 'count', 'find',
    'printf', 'scanf', 'main', 'tie',
}

# Tokenizer regex — match per token type
_TOKEN_RE = re.compile(
    r'(?P<comment>//[^\n]*)'
    r'|(?P<string>"(?:[^"\\]|\\.)*")'
    r'|(?P<char>\'(?:\\.|[^\'])\')'
    r'|(?P<preproc>\#\w+)'
    r'|(?P<number>\b\d+(?:\.\d+)?(?:LL|ULL|U|L)?\b)'
    r'|(?P<word>[A-Za-z_][A-Za-z_0-9]*)'
    r'|(?P<ws>[ \t]+)'
    r'|(?P<op>[<>+\-*/=!&|^%~?:.,;(){}\[\]])'
    r'|(?P<other>.)'
)

# Token color mapping
TOKEN_COLOR = {
    'comment': MUTED,
    'string': AMBER,
    'char': AMBER,
    'preproc': PINK,
    'number': ORANGE,
    'keyword': VIOLET,
    'type': CYAN,
    'builtin': PRIMARY_LT,
    'word': SILVER,
    'op': SILVER,
    'ws': SILVER,
    'other': SILVER,
}


def tokenize_cpp(line):
    """Tokenize satu baris C++ → list of (text, color)."""
    tokens = []
    for m in _TOKEN_RE.finditer(line):
        kind = m.lastgroup
        text = m.group()
        if kind == 'word':
            if text in CPP_KEYWORDS:
                color = TOKEN_COLOR['keyword']
            elif text in CPP_TYPES:
                color = TOKEN_COLOR['type']
            elif text in CPP_BUILTINS:
                color = TOKEN_COLOR['builtin']
            else:
                color = TOKEN_COLOR['word']
        else:
            color = TOKEN_COLOR.get(kind, SILVER)
        tokens.append((text, color))
    return tokens


# ═══════════════════════════════════════════════════════════════════
# 🎨 Code Block (with syntax highlighting)
# ═══════════════════════════════════════════════════════════════════

def code_box(slide, code_text, x, y, w, h, font_size=10, accent=PRIMARY,
             with_highlight=True):
    """Code block dengan syntax highlighting + modern dark style."""
    # Background dengan border accent
    rounded_rect(slide, x, y, w, h, BG_CODE,
                 line_color=accent, line_w=1.0)
    # Top "title bar" subtle (3 dot decoration like terminal)
    rect(slide, x, y, w, 0.22, BG_CARD)
    for i, col in enumerate([ROSE, AMBER, EMERALD]):
        circle(slide, x + 0.18 + i * 0.18, y + 0.11, 0.06, col)

    # Text content
    txBox = slide.shapes.add_textbox(
        Inches(x + 0.15), Inches(y + 0.3),
        Inches(w - 0.3), Inches(h - 0.4))
    tf = txBox.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)

    lines = code_text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT

        if not with_highlight:
            run = p.add_run()
            run.text = line
            run.font.size = Pt(font_size)
            run.font.color.rgb = SILVER
            run.font.name = "Menlo"
            continue

        if not line.strip():
            run = p.add_run()
            run.text = " "
            run.font.size = Pt(font_size)
            run.font.name = "Menlo"
            continue

        tokens = tokenize_cpp(line)
        for text, color in tokens:
            run = p.add_run()
            run.text = text
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.name = "Menlo"


def plain_code_box(slide, code_text, x, y, w, h, font_size=10, accent=PRIMARY,
                   text_color=SILVER):
    """Code box tanpa syntax highlighting (untuk pseudocode atau output)."""
    rounded_rect(slide, x, y, w, h, BG_CODE,
                 line_color=accent, line_w=1.0)
    rect(slide, x, y, w, 0.22, BG_CARD)
    for i, col in enumerate([ROSE, AMBER, EMERALD]):
        circle(slide, x + 0.18 + i * 0.18, y + 0.11, 0.06, col)

    txBox = slide.shapes.add_textbox(
        Inches(x + 0.15), Inches(y + 0.3),
        Inches(w - 0.3), Inches(h - 0.4))
    tf = txBox.text_frame
    tf.word_wrap = False
    lines = code_text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line if line.strip() else " "
        run.font.size = Pt(font_size)
        run.font.color.rgb = text_color
        run.font.name = "Menlo"


# ═══════════════════════════════════════════════════════════════════
# 🎨 Visual Components
# ═══════════════════════════════════════════════════════════════════

def card(slide, x, y, w, h, accent=PRIMARY, fill=None):
    if fill is None:
        fill = BG_CARD
    bg = rounded_rect(slide, x, y, w, h, fill,
                      line_color=BORDER, line_w=0.5)
    rect(slide, x, y, 0.08, h, accent)
    return bg


def info_card(slide, title, content_lines, x, y, w, h,
              accent=PRIMARY, title_size=12, body_size=10):
    """Card dengan title bar + bullet content."""
    bg = rounded_rect(slide, x, y, w, h, BG_CARD,
                      line_color=accent, line_w=1.2)
    title_h = 0.42
    rect(slide, x, y, w, title_h, accent)
    tb(slide, title, x + 0.18, y + 0.04, w - 0.3, title_h - 0.06,
       size=title_size, bold=True, color=BG_DEEP)
    cy = y + title_h + 0.12
    for line in content_lines:
        if isinstance(line, tuple):
            text, line_color = line[0], line[1]
        else:
            text, line_color = line, SILVER
        tb(slide, text, x + 0.22, cy, w - 0.4, 0.32,
           size=body_size, color=line_color)
        cy += 0.32
    return bg


def warn_card(slide, title, content_lines, x, y, w, h):
    return info_card(slide, "⚠ " + title, content_lines, x, y, w, h, accent=ROSE)


def tip_card(slide, title, content_lines, x, y, w, h):
    return info_card(slide, "✓ " + title, content_lines, x, y, w, h, accent=EMERALD)


def insight_card(slide, title, content_lines, x, y, w, h):
    return info_card(slide, "💡 " + title, content_lines, x, y, w, h, accent=AMBER)


def metric_card(slide, label, value, x, y, w, h, accent=PRIMARY):
    rounded_rect(slide, x, y, w, h, BG_CARD,
                 line_color=accent, line_w=1.0)
    rect(slide, x, y, w, 0.06, accent)
    tb(slide, label, x + 0.1, y + 0.18, w - 0.2, 0.4,
       size=10, color=SLATE, align=PP_ALIGN.CENTER)
    tb(slide, value, x + 0.1, y + 0.55, w - 0.2, h - 0.65,
       size=24, bold=True, color=accent, align=PP_ALIGN.CENTER)


def numbered_step_card(slide, num, title, desc, x, y, w, h,
                       accent=PRIMARY):
    rounded_rect(slide, x, y, w, h, BG_CARD,
                 line_color=accent, line_w=1.0)
    circle(slide, x + 0.55, y + 0.55, 0.36, accent)
    tb(slide, str(num), x + 0.18, y + 0.3, 0.75, 0.5,
       size=20, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)
    tb(slide, title, x + 1.15, y + 0.18, w - 1.25, 0.42,
       size=13, bold=True, color=accent)
    tb(slide, desc, x + 1.15, y + 0.6, w - 1.25, h - 0.7,
       size=10, color=SILVER)


def pill(slide, text, x, y, w, h, color=PRIMARY,
         text_color=BG_DEEP, size=9):
    rounded_rect(slide, x, y, w, h, color)
    tb(slide, text, x, y + h * 0.1, w, h * 0.85,
       size=size, bold=True, color=text_color, align=PP_ALIGN.CENTER)


def callout(slide, text, x, y, w, h, color=AMBER, size=12, icon=""):
    rounded_rect(slide, x, y, w, h, BG_CARD,
                 line_color=color, line_w=0.75)
    rect(slide, x, y, 0.1, h, color)
    if icon:
        tb(slide, icon, x + 0.2, y + 0.1, 0.5, h - 0.2,
           size=size + 4, color=color)
        tb(slide, text, x + 0.75, y + 0.12, w - 0.85, h - 0.24,
           size=size, color=WHITE)
    else:
        tb(slide, text, x + 0.25, y + 0.12, w - 0.4, h - 0.24,
           size=size, color=WHITE)


def stat_pill(slide, label, value, x, y, w, h, color=PRIMARY):
    """Pill horizontal: 'label: value'."""
    rounded_rect(slide, x, y, w, h, BG_CARD,
                 line_color=color, line_w=0.75)
    tb(slide, label, x + 0.15, y + 0.05, w * 0.45, h - 0.1,
       size=10, color=SLATE, align=PP_ALIGN.LEFT)
    tb(slide, value, x + w * 0.45, y + 0.05, w * 0.5 - 0.15, h - 0.1,
       size=12, bold=True, color=color, align=PP_ALIGN.RIGHT)


# ═══════════════════════════════════════════════════════════════════
# 📐 Standard Layouts
# ═══════════════════════════════════════════════════════════════════

def gradient_bg(slide):
    """Subtle gradient overlay — multiple colored rects for depth."""
    rect(slide, 0, 0, 13.33, 7.5, BG_MAIN)
    # Top-left glow
    circle(slide, -1, -1, 4, BG_CARD)
    # Bottom-right glow
    circle(slide, 14, 8, 5, BG_CARD)


def cover_slide(prs, day_num, title, subtitle_lines, decoration=None,
                topic_color=PRIMARY):
    slide = new_slide(prs, bg=BG_DEEP)
    # Layered background for depth
    rect(slide, 0, 0, 13.33, 7.5, BG_DEEP)
    rect(slide, 0, 0, 5.4, 7.5, BG_MAIN)
    rect(slide, 5.4, 0, 0.05, 7.5, topic_color)
    # Decorative geometric shapes
    circle(slide, 4.8, 0.5, 1.5, BG_CARD)
    circle(slide, -0.5, 7.5, 2.0, BG_CARD)
    # Main accent strip
    rect(slide, 0, 0, 0.3, 7.5, topic_color)
    rect(slide, 0.3, 0, 0.04, 7.5, CYAN)

    # Day badge
    rounded_rect(slide, 0.6, 0.5, 1.6, 0.45, topic_color)
    tb(slide, "DAY", 0.6, 0.55, 1.6, 0.35,
       size=14, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)

    # Big day number
    tb(slide, str(day_num).zfill(2), 0.6, 1.4, 4.8, 2.6,
       size=140, bold=True, color=topic_color)

    # Decorative lines
    rect(slide, 0.6, 4.0, 4.0, 0.08, CYAN)
    rect(slide, 0.6, 4.0, 1.5, 0.08, EMERALD)

    # Topic info
    tb(slide, "COMPETITIVE PROGRAMMING", 0.6, 4.2, 4.5, 0.32,
       size=11, bold=True, color=SLATE)
    tb(slide, "150 Day Roadmap", 0.6, 4.5, 4.5, 0.32,
       size=10, color=MUTED)

    # Title (right side)
    tb(slide, title, 6.0, 2.3, 7.0, 1.4,
       size=42, bold=True, color=WHITE)
    rect(slide, 6.0, 3.55, 0.5, 0.08, topic_color)
    rect(slide, 6.6, 3.55, 4.0, 0.08, CYAN)
    for i, line in enumerate(subtitle_lines):
        col = CYAN if i % 2 == 0 else VIOLET
        tb(slide, line, 6.0, 3.75 + i * 0.42, 7.0, 0.4,
           size=15, color=col)

    # Footer
    rect(slide, 6.0, 6.7, 0.4, 0.05, topic_color)
    tb(slide, "Modern Cobalt Theme · python-pptx",
       6.0, 6.8, 7.0, 0.3, size=9, color=MUTED, italic=True)

    if decoration:
        decoration(slide)
    return slide


def section_divider(prs, num, title, subtitle="", color=PRIMARY):
    slide = new_slide(prs, bg=BG_DEEP)
    rect(slide, 0, 0, 13.33, 7.5, BG_DEEP)
    # Big colored panel left
    rect(slide, 0, 0, 6.0, 7.5, BG_MAIN)
    rect(slide, 0, 0, 0.4, 7.5, color)
    rect(slide, 0.4, 0, 0.08, 7.5, CYAN)

    # Decorative geometric circles
    circle(slide, 11.5, 1.5, 1.0, color)
    circle(slide, 12.8, 2.5, 0.4, CYAN)
    circle(slide, 11.0, 3.0, 0.25, VIOLET)
    circle(slide, 12.0, 0.8, 0.18, EMERALD)
    circle(slide, 12.5, 6.5, 0.6, BG_CARD_2)
    circle(slide, 10.5, 7.0, 0.35, color)

    # Section label
    rect(slide, 1.0, 1.6, 1.5, 0.4, color)
    tb(slide, "SECTION", 1.0, 1.65, 1.5, 0.3,
       size=12, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)

    # Big number
    tb(slide, num, 1.0, 2.2, 5.0, 2.4,
       size=180, bold=True, color=color)

    # Title
    rect(slide, 1.0, 4.7, 0.5, 0.1, color)
    rect(slide, 1.6, 4.7, 5.0, 0.1, CYAN)
    tb(slide, title, 1.0, 4.85, 11.0, 1.0,
       size=44, bold=True, color=WHITE)
    if subtitle:
        tb(slide, subtitle, 1.0, 5.85, 11.0, 0.6,
           size=16, color=SLATE, italic=True)
    return slide


def slide_header(slide, title, subtitle="", color=PRIMARY):
    tb(slide, title, 0.5, 0.25, 12.3, 0.55,
       size=24, bold=True, color=WHITE)
    rect(slide, 0.5, 0.78, 0.4, 0.06, color)
    rect(slide, 0.95, 0.78, 11.85, 0.06, BORDER)
    if subtitle:
        tb(slide, subtitle, 0.5, 0.85, 12.3, 0.32,
           size=11, color=SLATE, italic=True)


def agenda_slide(prs, title, items):
    """4 kolom kartu vertikal dengan number, label, color, bullets."""
    slide = new_slide(prs)
    slide_header(slide, title, "Topik utama yang akan kita bahas hari ini")

    col_w = 3.05
    gap = 0.13
    for ci, (num, label, color, bullets) in enumerate(items):
        cx = 0.5 + ci * (col_w + gap)
        # Card background
        rounded_rect(slide, cx, 1.35, col_w, 5.7, BG_CARD,
                     line_color=color, line_w=1.0)
        rect(slide, cx, 1.35, col_w, 0.1, color)
        # Number with bg
        rounded_rect(slide, cx + 0.2, 1.6, 1.0, 0.6, color)
        tb(slide, num, cx + 0.2, 1.65, 1.0, 0.5,
           size=22, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)
        # Label
        tb(slide, label, cx + 0.2, 2.4, col_w - 0.4, 0.85,
           size=15, bold=True, color=WHITE)
        # Divider
        rect(slide, cx + 0.2, 3.25, col_w - 0.4, 0.04, color)
        # Bullets
        for ii, item in enumerate(bullets):
            tb(slide, "▸  " + item, cx + 0.2, 3.4 + ii * 0.5,
               col_w - 0.4, 0.45, size=10.5, color=SILVER)
    return slide


def closing_slide(prs, day_num, title_main, takeaways):
    slide = new_slide(prs, bg=BG_DEEP)
    rect(slide, 0, 0, 13.33, 7.5, BG_DEEP)
    rect(slide, 0, 0, 0.3, 7.5, PRIMARY)
    rect(slide, 0.3, 0, 0.06, 7.5, CYAN)
    # Decorative
    circle(slide, 12.5, 1.0, 1.0, BG_CARD)
    circle(slide, 13.5, 6.5, 1.5, BG_CARD)

    # Title
    rounded_rect(slide, 0.6, 0.4, 5.5, 0.7, PRIMARY)
    tb(slide, "🎯 Key Takeaways", 0.6, 0.45, 5.5, 0.6,
       size=24, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)
    tb(slide, title_main, 0.6, 1.15, 12.0, 0.4,
       size=14, color=SLATE, italic=True)
    rect(slide, 0.6, 1.55, 12.0, 0.04, BORDER)

    for i, (col, title, desc) in enumerate(takeaways):
        ty = 1.85 + i * 1.0
        # Card
        rounded_rect(slide, 0.6, ty, 12.0, 0.85, BG_CARD,
                     line_color=col, line_w=0.75)
        rect(slide, 0.6, ty, 0.1, 0.85, col)
        # Title
        tb(slide, title, 0.85, ty + 0.1, 4.0, 0.4,
           size=14, bold=True, color=col)
        # Description
        tb(slide, desc, 0.85, ty + 0.45, 11.0, 0.4,
           size=10.5, color=SILVER)

    rect(slide, 0.6, 7.1, 12.0, 0.04, PRIMARY)
    tb(slide, f"Day {day_num} | Competitive Programming · 150 Day Roadmap",
       0.6, 7.18, 12.0, 0.3, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    return slide


# ═══════════════════════════════════════════════════════════════════
# 📊 Tables & Comparison
# ═══════════════════════════════════════════════════════════════════

def styled_table(slide, headers, rows, x, y, col_widths, row_h=0.5,
                 header_colors=None, body_size=10, header_size=11):
    if header_colors is None:
        header_colors = [PRIMARY] * len(headers)
    total_w = sum(col_widths)
    col_x = [x]
    for cw in col_widths[:-1]:
        col_x.append(col_x[-1] + cw)

    header_h = row_h * 0.85
    for ci, (h, col) in enumerate(zip(headers, header_colors)):
        rounded_rect(slide, col_x[ci], y, col_widths[ci] - 0.04, header_h, col)
        tb(slide, h, col_x[ci] + 0.1, y + 0.05,
           col_widths[ci] - 0.2, header_h - 0.1,
           size=header_size, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)

    for ri, row in enumerate(rows):
        ry = y + header_h + 0.05 + ri * row_h
        bg = BG_CARD if ri % 2 == 0 else BG_CARD_2
        rect(slide, x, ry, total_w - 0.04, row_h - 0.04, bg)
        for ci, cell in enumerate(row):
            text_color = WHITE
            if isinstance(cell, tuple):
                text, text_color = cell[0], cell[1]
            else:
                text = cell
            tb(slide, str(text), col_x[ci] + 0.1, ry + 0.08,
               col_widths[ci] - 0.2, row_h - 0.16,
               size=body_size, color=text_color,
               align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)


def comparison_panel(slide, x, y, w, h, left_title, left_lines,
                     right_title, right_lines,
                     left_color=ROSE, right_color=EMERALD):
    half_w = (w - 0.2) / 2
    info_card(slide, left_title, left_lines, x, y, half_w, h, accent=left_color)
    info_card(slide, right_title, right_lines, x + half_w + 0.2, y,
              half_w, h, accent=right_color)


def vs_divider(slide, x, y, h):
    cy = y + h / 2
    circle(slide, x, cy, 0.32, AMBER)
    tb(slide, "VS", x - 0.32, cy - 0.16, 0.64, 0.32,
       size=14, bold=True, color=BG_DEEP, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════
# 🎬 Step-by-step / Trace Diagrams
# ═══════════════════════════════════════════════════════════════════

def array_cells(slide, values, x, y, cell_w=0.55, cell_h=0.5,
                highlights=None, colors_map=None,
                show_index=True, value_size=14):
    if highlights is None:
        highlights = {}
    if colors_map is None:
        colors_map = {}
    for i, val in enumerate(values):
        cx = x + i * (cell_w + 0.05)
        cell_color = highlights.get(i, BG_CARD_2)
        text_color = colors_map.get(
            i, BG_DEEP if i in highlights else WHITE)
        rounded_rect(slide, cx, y, cell_w, cell_h, cell_color,
                     line_color=BORDER, line_w=0.5)
        if show_index:
            tb(slide, str(i), cx, y - 0.27, cell_w, 0.22,
               size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        tb(slide, str(val), cx + 0.02, y + cell_h * 0.18,
           cell_w - 0.04, cell_h * 0.7,
           size=value_size, bold=True, color=text_color,
           align=PP_ALIGN.CENTER)


def pointer_label(slide, label, x, y, color=PRIMARY, size=11):
    tb(slide, "↑", x, y, 0.4, 0.25,
       size=14, bold=True, color=color, align=PP_ALIGN.CENTER)
    tb(slide, label, x, y + 0.2, 0.4, 0.3,
       size=size, bold=True, color=color, align=PP_ALIGN.CENTER)


def step_track(slide, x, y, total_steps, current_step, color=PRIMARY):
    dot_r = 0.08
    spacing = 0.4
    for i in range(total_steps):
        dx = x + i * spacing
        col = color if i <= current_step else DIM
        circle(slide, dx, y, dot_r, col)
        if i < total_steps - 1:
            rect(slide, dx + dot_r, y - 0.02, spacing - 2 * dot_r, 0.04,
                 col if i < current_step else DIM)


def flow_node(slide, label, x, y, w, h, color=PRIMARY, text_color=BG_DEEP):
    rounded_rect(slide, x, y, w, h, color)
    tb(slide, label, x + 0.05, y + 0.05, w - 0.1, h - 0.1,
       size=11, bold=True, color=text_color, align=PP_ALIGN.CENTER)


def flow_arrow(slide, x1, y, x2, color=PRIMARY, thickness=0.04):
    rect(slide, x1, y, x2 - x1 - 0.18, thickness, color)
    shape = slide.shapes.add_shape(
        5, Inches(x2 - 0.2), Inches(y - 0.08),
        Inches(0.22), Inches(0.2 + thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


# ═══════════════════════════════════════════════════════════════════
# � Lines, Edges, & Connectors — proper diagonal lines
# ═══════════════════════════════════════════════════════════════════

def line(slide, x1, y1, x2, y2, color=BORDER, thickness=1.0):
    """Garis lurus dari (x1,y1) ke (x2,y2) — dalam inches.
    Pakai connector untuk akurasi dan dukungan diagonal."""
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(thickness)
    return cn


def arrow(slide, x1, y1, x2, y2, color=PRIMARY, thickness=1.5,
          head_size="medium"):
    """Garis dengan arrow head di ujung (x2,y2)."""
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(thickness)
    # Set arrow head di end (line element XML)
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = cn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    return cn


def dashed_line(slide, x1, y1, x2, y2, color=BORDER, thickness=1.0):
    """Garis dengan style dashed."""
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(thickness)
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = cn.line._get_or_add_ln()
    prst = etree.SubElement(ln, qn('a:prstDash'))
    prst.set('val', 'dash')
    return cn


# ═══════════════════════════════════════════════════════════════════
# 🌳 Tree Drawing Helpers
# ═══════════════════════════════════════════════════════════════════

def tree_node(slide, cx, cy, label, color=PRIMARY, radius=0.28,
              text_color=None, size=12):
    """Lingkaran node tree dengan label di tengah."""
    if text_color is None:
        text_color = BG_DEEP
    circle(slide, cx, cy, radius, color)
    tb(slide, str(label),
       cx - radius, cy - radius * 0.55,
       2 * radius, radius * 1.1,
       size=size, bold=True, color=text_color, align=PP_ALIGN.CENTER)


def tree_edge(slide, x1, y1, x2, y2, r1=0.28, r2=0.28,
              color=SLATE, thickness=1.5):
    """Garis tree dari node (x1,y1) radius r1 → node (x2,y2) radius r2.
    Otomatis menyesuaikan agar tidak menumpuk node."""
    dx = x2 - x1
    dy = y2 - y1
    d = math.hypot(dx, dy)
    if d == 0:
        return
    ux, uy = dx / d, dy / d
    sx = x1 + ux * r1
    sy = y1 + uy * r1
    ex = x2 - ux * r2
    ey = y2 - uy * r2
    line(slide, sx, sy, ex, ey, color=color, thickness=thickness)


def draw_tree(slide, nodes, edges, default_color=PRIMARY,
              edge_color=SLATE, radius=0.28, value_size=12):
    """Gambar tree berdasarkan nodes dan edges.
    nodes: list of (label, x, y) atau (label, x, y, color).
    edges: list of (idx_from, idx_to) atau (idx_from, idx_to, edge_color).
    """
    # Draw edges first (di belakang nodes)
    for e in edges:
        if len(e) == 3:
            i, j, ec = e
        else:
            i, j = e
            ec = edge_color
        x1, y1 = nodes[i][1], nodes[i][2]
        x2, y2 = nodes[j][1], nodes[j][2]
        r1 = nodes[i][4] if len(nodes[i]) > 4 else radius
        r2 = nodes[j][4] if len(nodes[j]) > 4 else radius
        tree_edge(slide, x1, y1, x2, y2, r1=r1, r2=r2,
                  color=ec, thickness=1.75)
    # Draw nodes on top
    for n in nodes:
        if len(n) >= 4:
            label, x, y, col = n[0], n[1], n[2], n[3]
        else:
            label, x, y = n[0], n[1], n[2]
            col = default_color
        r = n[4] if len(n) > 4 else radius
        tree_node(slide, x, y, label, color=col, radius=r,
                  size=value_size)


# ═══════════════════════════════════════════════════════════════════
# 🕸️ Graph Drawing Helpers
# ═══════════════════════════════════════════════════════════════════

def graph_node(slide, cx, cy, label, color=PRIMARY, radius=0.32,
               text_color=None, size=14):
    """Lingkaran graph node (sedikit lebih besar dari tree)."""
    if text_color is None:
        text_color = BG_DEEP
    circle(slide, cx, cy, radius, color, line_color=WHITE, line_w=1.0)
    tb(slide, str(label),
       cx - radius, cy - radius * 0.5,
       2 * radius, radius,
       size=size, bold=True, color=text_color, align=PP_ALIGN.CENTER)


def draw_graph(slide, nodes, edges, default_color=PRIMARY,
               edge_color=SLATE, edge_thickness=1.75,
               radius=0.32, value_size=14, edge_labels=None):
    """Gambar graph dengan nodes dan edges.
    nodes: list of (label, x, y) atau (label, x, y, color).
    edges: list of (i, j) atau (i, j, edge_color).
    edge_labels: dict {(i, j): "weight_label"} optional.
    """
    if edge_labels is None:
        edge_labels = {}
    # Edges first
    for e in edges:
        if len(e) == 3:
            i, j, ec = e
        else:
            i, j = e
            ec = edge_color
        x1, y1 = nodes[i][1], nodes[i][2]
        x2, y2 = nodes[j][1], nodes[j][2]
        r1 = nodes[i][4] if len(nodes[i]) > 4 else radius
        r2 = nodes[j][4] if len(nodes[j]) > 4 else radius
        tree_edge(slide, x1, y1, x2, y2, r1=r1, r2=r2,
                  color=ec, thickness=edge_thickness)
        # Edge label di tengah
        key = (i, j) if (i, j) in edge_labels else (j, i)
        if key in edge_labels:
            mx = (x1 + x2) / 2
            my = (y1 + y2) / 2
            lbl = edge_labels[key]
            rounded_rect(slide, mx - 0.18, my - 0.12, 0.36, 0.24,
                         BG_DEEP, line_color=ec, line_w=0.5)
            tb(slide, str(lbl), mx - 0.18, my - 0.1, 0.36, 0.2,
               size=9, bold=True, color=ec, align=PP_ALIGN.CENTER)
    # Nodes
    for n in nodes:
        if len(n) >= 4:
            label, x, y, col = n[0], n[1], n[2], n[3]
        else:
            label, x, y = n[0], n[1], n[2]
            col = default_color
        r = n[4] if len(n) > 4 else radius
        graph_node(slide, x, y, label, color=col, radius=r,
                   size=value_size)


def directed_edge(slide, x1, y1, x2, y2, r1=0.32, r2=0.32,
                  color=PRIMARY, thickness=1.75):
    """Edge dengan arrow head — menyesuaikan radius node."""
    dx = x2 - x1
    dy = y2 - y1
    d = math.hypot(dx, dy)
    if d == 0:
        return
    ux, uy = dx / d, dy / d
    sx = x1 + ux * r1
    sy = y1 + uy * r1
    ex = x2 - ux * r2
    ey = y2 - uy * r2
    arrow(slide, sx, sy, ex, ey, color=color, thickness=thickness)


# ═══════════════════════════════════════════════════════════════════
# �🚀 Helper Builders
# ═══════════════════════════════════════════════════════════════════

def make_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def save_presentation(prs, file_path):
    prs.save(file_path)
    print(f"✓ Saved: {file_path}")
